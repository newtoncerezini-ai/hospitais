from __future__ import annotations

import argparse
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = ROOT / "data" / "raw" / "health-units.xlsx"
DEFAULT_MAP_DATA = ROOT / "data" / "reference" / "pernambuco-map.json"
DEFAULT_OUTPUT = ROOT / "public" / "data" / "health-units.json"
DEFAULT_QUALITY_OUTPUT = ROOT / "data" / "processed" / "data-quality.json"
SHEET_NAME = "Consolidado"
DRIVE_SHEET_NAME = "NÃO MEXER Consolidado"
CONSTRUCTION_SHEET_NAME = "UNIDADES EM CONSTRUÇÃO"
DRIVE_SOURCE_SHEETS = (
    "Hospitais Regionais e OSS",
    "Os Seis Grandes",
    "UPA e UPA-E",
    "Rede Credenciada",
    CONSTRUCTION_SHEET_NAME,
)


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def normalize(value: Any) -> str:
    text = clean_text(value) or ""
    ascii_text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "", ascii_text.lower())


def slugify(value: str) -> str:
    ascii_text = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^a-z0-9]+", "-", ascii_text.lower()).strip("-")


def title_particles(value: str) -> str:
    result = value
    for particle in (" Do ", " Da ", " Dos ", " Das ", " De "):
        result = result.replace(particle, particle.lower())
    return result


def normalize_status(value: Any) -> str:
    status = clean_text(value)
    normalized_status = normalize(status)
    if "construcao" in normalized_status:
        return "Em construção"
    if "funcionamento" in normalized_status:
        return "Em funcionamento"
    return status or "Não informado"


def normalize_unit_type(value: Any) -> str:
    unit_type = clean_text(value) or "Não informado"
    normalized_type = normalize(unit_type)
    if normalized_type == "hospital":
        return "Hospital"
    if normalized_type in {"upaer", "upaerregional"}:
        return "UPAE-R"
    if normalized_type in {"grandeemergencia", "grandesemergencias"}:
        return "Hospital"
    return unit_type


def parse_integer(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (date, datetime)):
        return (value.date() - date(1899, 12, 30)).days if isinstance(value, datetime) else (value - date(1899, 12, 30)).days
    if isinstance(value, (int, float)):
        return int(value)
    text = clean_text(value)
    if not text or text in {"-", "—"}:
        return None
    match = re.search(r"\d+", text.replace(".", ""))
    return int(match.group()) if match else None


def parse_money(value: Any) -> dict[str, Any]:
    text = clean_text(value)
    if isinstance(value, bool) or value is None:
        return {"amount": None, "label": None}
    if isinstance(value, (int, float)):
        return {"amount": round(float(value), 2), "label": None}
    if not text or text in {"-", "—"}:
        return {"amount": None, "label": None}

    normalized_text = normalize(text)
    multiplier = 1
    if "milhao" in normalized_text or "milhoes" in normalized_text:
        multiplier = 1_000_000
    elif re.search(r"\bmil\b", text.lower()):
        multiplier = 1_000

    compact = re.sub(r"[^0-9,.]", "", text)
    decimal_match = re.search(r"([,.])(\d{2})$", compact)
    amount = None
    if multiplier > 1:
        match = re.search(r"\d+(?:[.,]\d+)?", text)
        if match:
            amount = float(match.group().replace(".", "").replace(",", ".")) * multiplier
    elif decimal_match:
        cents = decimal_match.group(2)
        whole = re.sub(r"\D", "", compact[: decimal_match.start()]) or "0"
        amount = float(f"{whole}.{cents}")
    elif compact and re.search(r"\d", compact):
        amount = float(re.sub(r"\D", "", compact))

    return {"amount": round(amount, 2) if amount is not None else None, "label": text}


def parse_percentage(value: Any) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 10)
    text = clean_text(value)
    if not text or text in {"-", "—"} or text.startswith("#"):
        return None
    match = re.search(r"-?\d+(?:[.,]\d+)?", text.replace(".", ""))
    if not match:
        return None
    parsed = float(match.group().replace(",", "."))
    return round(parsed / 100 if "%" in text else parsed, 10)


def normalize_ibge(value: Any) -> str | None:
    if value is None or isinstance(value, bool):
        return None
    try:
        return str(int(float(value))).zfill(7)
    except (TypeError, ValueError):
        digits = re.sub(r"\D", "", str(value))
        return digits.zfill(7) if digits else None


def has_usable_value(value: Any) -> bool:
    text = clean_text(value)
    return bool(
        text
        and text not in {"-", "—"}
        and not text.startswith("#")
        and normalize(text) not in {"aguardandoinformacao", "naoinformado", "naoencontrado"}
    )


def canonical_unit_key(value: Any) -> str:
    key = normalize(value)
    aliases = {
        "hospitalprofagamenonmagalhaes": "hospitalprofessoragamenonmagalhaes",
        "maternidadebritesdealbuquerque": "hospitalbritesdealbuquerque",
        "upaerdorecife": "upaereabilitacaodorecife",
    }
    return aliases.get(key, key)


def find_header_row(sheet: Any, required_headers: set[str], max_rows: int = 12) -> int:
    for row_number in range(1, min(sheet.max_row or max_rows, max_rows) + 1):
        normalized = {normalize(cell.value) for cell in sheet[row_number] if clean_text(cell.value)}
        if normalized & required_headers:
            return row_number
    raise ValueError(f"Cabeçalho não encontrado na aba {sheet.title}")


def read_named_table(sheet: Any, unit_headers: tuple[str, ...]) -> tuple[list[str], list[dict[str, Any]]]:
    normalized_unit_headers = {normalize(header) for header in unit_headers}
    header_row = find_header_row(sheet, normalized_unit_headers)
    raw_headers = [clean_text(cell.value) or "" for cell in sheet[header_row]]
    normalized_headers = [normalize(header) for header in raw_headers]
    rows: list[dict[str, Any]] = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        row: dict[str, Any] = {
            "__source_row": source_row,
            "__source_sheet": sheet.title,
        }
        for key, value in zip(normalized_headers, values):
            if not key:
                continue
            if key not in row or not has_usable_value(row[key]):
                row[key] = value
        unit_name = next((clean_text(row.get(key)) for key in normalized_unit_headers if clean_text(row.get(key))), None)
        if unit_name:
            row["__unit_name"] = unit_name
            rows.append(row)
    return raw_headers, rows


def first_value(row: dict[str, Any] | None, *headers: str) -> Any:
    if not row:
        return None
    for header in headers:
        value = row.get(normalize(header))
        if has_usable_value(value):
            return value
    return None


def text_field(text: Any, label: str) -> str | None:
    target = normalize(label)
    for line in str(text or "").splitlines():
        if ":" not in line:
            continue
        field, value = line.split(":", 1)
        if normalize(field) == target and has_usable_value(value):
            return clean_text(value)
    return None


def load_previous_units() -> dict[str, dict[str, Any]]:
    if not DEFAULT_OUTPUT.exists():
        return {}
    try:
        payload = json.loads(DEFAULT_OUTPUT.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    previous: dict[str, dict[str, Any]] = {}
    for unit in payload.get("units", []):
        previous.setdefault(canonical_unit_key(unit.get("name")), unit)
    return previous


def inferred_type(unit_name: str, raw_type: Any, previous: dict[str, Any] | None) -> str:
    normalized_type = normalize(raw_type)
    normalized_name = normalize(unit_name)
    if "construcao" in normalized_type:
        return "UPAE-R" if normalized_name.startswith("upaer") else "Hospital"
    if normalized_type in {"upaeupa", "upaupae"}:
        return previous.get("type", "UPAE") if previous else "UPAE"
    return normalize_unit_type(raw_type)


def build_drive_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    consolidated = workbook[DRIVE_SHEET_NAME]
    headers, consolidated_rows = read_named_table(consolidated, ("Unidade",))
    previous_units = load_previous_units()

    supplements: dict[str, dict[str, Any]] = {}
    source_specs = {
        "Hospitais Regionais e OSS": ("HOSPITAL",),
        "Os Seis Grandes": ("HOSPITAL",),
        "UPA e UPA-E": ("UPA ou UPAE",),
        CONSTRUCTION_SHEET_NAME: ("Novas Maternidades",),
    }
    for sheet_name, unit_headers in source_specs.items():
        _, source_rows = read_named_table(workbook[sheet_name], unit_headers)
        for source_row in source_rows:
            supplements[canonical_unit_key(source_row["__unit_name"])] = source_row

    canonical_rows: list[dict[str, Any]] = []
    for source_row in consolidated_rows:
        unit_name = clean_text(source_row["__unit_name"]) or "Unidade sem nome"
        key = canonical_unit_key(unit_name)
        supplement = supplements.get(key)
        previous = previous_units.get(key)
        raw_type = first_value(source_row, "Tipo")
        status = "Em construção" if "construcao" in normalize(raw_type) else "Em funcionamento"
        source_text = first_value(source_row, "Texto")
        management_type = first_value(
            source_row,
            "Tipo de Gestão",
        ) or first_value(
            supplement,
            "Tipo de Gestão (OSS / Gestão Própria)",
            "Tipo de Gestão",
        ) or (previous or {}).get("managementType")
        operator = first_value(supplement, "OSS")
        inherited_fields: list[str] = []
        if not operator and normalize(management_type) == "gestaopropria":
            operator = "SES"
        if not operator and previous and previous.get("management"):
            operator = previous["management"]
            inherited_fields.append("management")

        address = first_value(supplement, "Localização") or text_field(source_text, "Localização")
        if not address and previous and previous.get("address"):
            address = previous["address"]
            inherited_fields.append("address")

        canonical_rows.append(
            {
                "NOME MUNICIPIO": first_value(source_row, "MUNICIPIO") or (previous or {}).get("municipality"),
                "RD": first_value(source_row, "RD") or (previous or {}).get("rd"),
                "GERES": first_value(source_row, "GERES") or first_value(supplement, "GERES"),
                "UNIDADE DE SAÚDE": unit_name,
                "LOCALIZAÇÃO": address,
                "STATUS": status,
                "TIPO": inferred_type(unit_name, raw_type, previous),
                "TIPO GESTÃO": management_type,
                "GESTÃO": operator,
                "LEITOS": first_value(source_row, "Nº de Leitos") or first_value(supplement, "Número de leitos", "Leitos", "Leitos previstos"),
                "LEITOS ABERTOS NESTA GESTÃO": first_value(supplement, "Leitos abertos nesta gestão"),
                "TIPOS DE LEITOS ABERTOS NESTA GESTÃO": first_value(supplement, "Tipos de leitos abertos nesta gestão"),
                "LEITOS A ABRIR COM O FIM DA REFORMA": first_value(supplement, "Leitos a abrir com o fim das reformas"),
                "PERFIL": first_value(source_row, "Perfil") or first_value(supplement, "Perfil"),
                "PROSSIONAIS": first_value(source_row, "Nº Profissionais Total") or first_value(supplement, "Nº Profissionais Total"),
                "PROFISSIONAIS CONVOCADOS NA GESTÃO": first_value(source_row, "Nº Profissionais convocados nesta gestão") or first_value(supplement, "Nº Profissionais convocados nesta gestão"),
                "CONTRATO DE MANUTENÇÃO PREDIAL": first_value(source_row, "Valor do contrato de manutenção predial (Anual)") or first_value(supplement, "Valor do contrato de manutenção predial (anual)"),
                "INVESTIMENTO NESTA GESTÃO": first_value(source_row, "Investimento nesta gestão") or first_value(supplement, "Investimento nesta gestão¹", "Investimento nesta gestão"),
                "PRINCIPAIS AVANÇOS": first_value(source_row, "Principais avanços nessa gestão") or first_value(supplement, "Principais avanços nessa gestão"),
                "COFINANCIAMENTO 2022": first_value(source_row, "Cofinanciamento 2022 (Apenas Gestão Própria)") or first_value(supplement, "Cofinanciamento 2022 (Apenas Gestão Própria)", "Cofinanciamento 2022"),
                "COFINANCIAMENTO 2025": first_value(source_row, "Cofinanciamento 2025 (Apenas Gestão Própria)") or first_value(supplement, "Cofinanciamento 2025 (Apenas Gestão Própria)", "Cofinanciamento 2025"),
                "AUMENTO COFINANCIAMENTO": first_value(source_row, "Aumento Cofinanciamento (2022 a 2025) % (Apenas Gestão Própria)") or first_value(supplement, "Aumento Cofinanciamento (2022 a 2025) % (Apenas Gestão Própria)", "Aumento Cofinanciamento (2022 a 2025) %"),
                "REPASSE OS 2025": first_value(source_row, "Repasse (OS) em 2025") or first_value(supplement, "Repasse (OS) em 2025"),
                "INVESTIMENTO OBRA E EQUIPAGEM": first_value(source_row, "Investimento para obra e equipagem") or first_value(supplement, "Investimento para obra e equipagem"),
                "PROFISSIONAIS SERVIDORES": first_value(supplement, "Profissionais Servidores"),
                "PROFISSIONAIS COMISSIONADOS": first_value(supplement, "Comissionados"),
                "PROFISSIONAIS CLT": first_value(supplement, "Profissionais CLT"),
                "PROFISSIONAIS PJ": first_value(supplement, "Profissionais PJ"),
                "PROFISSIONAIS TERCEIRIZADOS": first_value(supplement, "Profissionais Terceirizados"),
                "JOVEM APRENDIZ": first_value(supplement, "Jovem Aprendiz"),
                "CONTRATOS TEMPORÁRIOS": first_value(supplement, "Contratos Temporários"),
                "CEDIDOS DE OUTROS ÓRGÃOS": first_value(supplement, "Cedidos de outros órgãos"),
                "TEXTO CONSOLIDADO": source_text,
                "__ibge_code": normalize_ibge(first_value(source_row, "CÓDIGO DO IBGE")),
                "__source_row": source_row["__source_row"],
                "__source_sheet": DRIVE_SHEET_NAME,
                "__reference_sources": [
                    {"sheet": supplement["__source_sheet"], "row": supplement["__source_row"]}
                ] if supplement else [],
                "__inherited_fields": inherited_fields,
            }
        )

    _, credentialed_rows = read_named_table(workbook["Rede Credenciada"], ("HOSPITAL",))
    for credentialed_row in credentialed_rows:
        unit_name = clean_text(credentialed_row["__unit_name"]) or "Unidade sem nome"
        previous = previous_units.get(canonical_unit_key(unit_name))
        municipality = first_value(credentialed_row, "Localização") or (previous or {}).get("municipality")
        canonical_rows.append(
            {
                "NOME MUNICIPIO": municipality,
                "RD": (previous or {}).get("rd"),
                "GERES": first_value(credentialed_row, "GERES") or (previous or {}).get("geres"),
                "UNIDADE DE SAÚDE": unit_name,
                "LOCALIZAÇÃO": None,
                "STATUS": None,
                "TIPO": "Rede Credenciada",
                "TIPO GESTÃO": None,
                "GESTÃO": None,
                "LEITOS": None,
                "LEITOS ABERTOS NESTA GESTÃO": first_value(credentialed_row, "Leitos abertos nesta gestão"),
                "TIPOS DE LEITOS ABERTOS NESTA GESTÃO": first_value(credentialed_row, "Tipos de leitos abertos nesta gestão"),
                "LEITOS A ABRIR COM O FIM DA REFORMA": None,
                "PERFIL": None,
                "PROSSIONAIS": None,
                "PROFISSIONAIS CONVOCADOS NA GESTÃO": None,
                "CONTRATO DE MANUTENÇÃO PREDIAL": None,
                "INVESTIMENTO NESTA GESTÃO": None,
                "PRINCIPAIS AVANÇOS": None,
                "__ibge_code": (previous or {}).get("ibgeCode"),
                "__source_row": credentialed_row["__source_row"],
                "__source_sheet": "Rede Credenciada",
                "__reference_sources": [],
                "__inherited_fields": ["municipality", "rd", "geres", "ibgeCode"] if previous else [],
            }
        )
    workbook.close()
    return headers, canonical_rows


def scan_drive_formula_errors(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=False, data_only=True)
    if DRIVE_SHEET_NAME not in workbook.sheetnames:
        workbook.close()
        return []
    sheet = workbook[DRIVE_SHEET_NAME]
    header_row = find_header_row(sheet, {normalize("Unidade")})
    errors: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, (sheet.max_row or header_row) + 1):
        unit_name = clean_text(sheet.cell(row_number, 5).value)
        if not unit_name:
            continue
        for column in range(1, (sheet.max_column or 20) + 1):
            value = sheet.cell(row_number, column).value
            if isinstance(value, str) and value.startswith("#"):
                errors.append({
                    "sheet": DRIVE_SHEET_NAME,
                    "cell": sheet.cell(row_number, column).coordinate,
                    "unitName": unit_name,
                    "error": value,
                })
    workbook.close()
    return errors


def read_rows(workbook_path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if DRIVE_SHEET_NAME in workbook.sheetnames:
        workbook.close()
        return build_drive_rows(workbook_path)
    if SHEET_NAME not in workbook.sheetnames:
        raise ValueError(f"Aba obrigatória ausente: {SHEET_NAME} ou {DRIVE_SHEET_NAME}")
    sheet = workbook[SHEET_NAME]
    iterator = sheet.iter_rows(values_only=True)
    headers = [clean_text(value) or "" for value in next(iterator)]
    rows: list[dict[str, Any]] = []
    for source_row, values in enumerate(iterator, start=2):
        row = dict(zip(headers, values))
        if clean_text(row.get("UNIDADE DE SAÚDE")):
            row["__source_row"] = source_row
            row["__source_sheet"] = SHEET_NAME
            rows.append(row)
    workbook.close()
    return headers, rows


def infer_construction_municipality(unit_name: str, location: Any) -> str | None:
    if normalize(unit_name).startswith("maternidadede"):
        return clean_text(re.sub(r"^Maternidade\s+de\s+", "", unit_name, flags=re.IGNORECASE))
    if normalize(unit_name).startswith("upaer"):
        return clean_text(re.sub(r"^UPAE[-\s]*R\s+", "", unit_name, flags=re.IGNORECASE))
    location_text = clean_text(location)
    if location_text not in {None, "-", "—"}:
        return location_text
    return None


def read_construction_rows(
    workbook_path: Path,
    municipality_reference: dict[str, dict[str, str]],
) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[CONSTRUCTION_SHEET_NAME]
    rows: list[dict[str, Any]] = []
    for source_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        unit_name = clean_text(values[1] if len(values) > 1 else None)
        if not unit_name:
            continue
        municipality = infer_construction_municipality(
            unit_name,
            values[2] if len(values) > 2 else None,
        ) or "Não informado"
        municipality_ref = municipality_reference.get(normalize(municipality))
        unit_type = "UPAE/R" if normalize(unit_name).startswith("upaer") else "Hospital"
        rows.append(
            {
                "NOME MUNICIPIO": municipality,
                "RD": municipality_ref["rd"] if municipality_ref else None,
                "GERES": values[3] if len(values) > 3 else None,
                "UNIDADE DE SAÚDE": unit_name,
                "LOCALIZAÇÃO": None,
                "STATUS": "Em construção",
                "TIPO": unit_type,
                "TIPO GESTÃO": None,
                "GESTÃO": None,
                "LEITOS": values[4] if len(values) > 4 else None,
                "PERFIL": values[5] if len(values) > 5 else None,
                "PROSSIONAIS": None,
                "PROFISSIONAIS CONVOCADOS NA GESTÃO": None,
                "CONTRATO DE MANUTENÇÃO PREDIAL": None,
                "INVESTIMENTO NESTA GESTÃO": values[6] if len(values) > 6 else None,
                "PRINCIPAIS AVANÇOS": None,
                "__source_row": source_row,
                "__source_sheet": CONSTRUCTION_SHEET_NAME,
            }
        )
    workbook.close()
    return rows


def merge_construction_rows(
    consolidated_rows: list[dict[str, Any]],
    construction_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    merged = list(consolidated_rows)
    by_name = {
        normalize(row.get("UNIDADE DE SAÚDE")): row
        for row in merged
        if clean_text(row.get("UNIDADE DE SAÚDE"))
    }
    for construction_row in construction_rows:
        key = normalize(construction_row.get("UNIDADE DE SAÚDE"))
        existing = by_name.get(key)
        if not existing:
            construction_row["__construction_source_row"] = construction_row["__source_row"]
            merged.append(construction_row)
            by_name[key] = construction_row
            continue
        if normalize_status(existing.get("STATUS")) == "Em funcionamento":
            continue
        existing["__construction_source_row"] = construction_row["__source_row"]
        for field in ("STATUS", "LEITOS", "PERFIL", "INVESTIMENTO NESTA GESTÃO"):
            if not clean_text(existing.get(field)) and clean_text(construction_row.get(field)):
                existing[field] = construction_row[field]
    return merged


def read_municipality_reference(workbook_path: Path) -> dict[str, dict[str, str]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook["0_Template_Cod_Mun_RD"]
    header_row = find_header_row(sheet, {normalize("NOME MUNICIPIO")})
    reference: dict[str, dict[str, str]] = {}
    for code_value, municipality_value, rd_value, *_ in sheet.iter_rows(min_row=header_row + 1, values_only=True):
        municipality = clean_text(municipality_value)
        rd = clean_text(rd_value)
        if not municipality or not rd or code_value is None:
            continue
        try:
            code = str(int(code_value)).zfill(7)
        except (TypeError, ValueError):
            continue
        reference[normalize(municipality)] = {
            "municipality": municipality,
            "code": code,
            "rd": title_particles(rd),
        }
    workbook.close()
    return reference


def read_source_investments(workbook_path: Path) -> dict[str, Any]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    values: dict[str, Any] = {}
    ignored_sheets = {SHEET_NAME, "0_Template_Cod_Mun_RD", CONSTRUCTION_SHEET_NAME}
    for sheet in workbook.worksheets:
        if sheet.title in ignored_sheets:
            continue
        rows = sheet.iter_rows(values_only=True)
        headers = [clean_text(value) or "" for value in next(rows)]
        unit_index = next((index for index, header in enumerate(headers) if header in {"HOSPITAL", "UPA ou UPAE"}), None)
        investment_index = next((index for index, header in enumerate(headers) if "Investimento nesta gestão" in header), None)
        if unit_index is None or investment_index is None:
            continue
        for row in rows:
            unit = clean_text(row[unit_index] if unit_index < len(row) else None)
            raw_investment = row[investment_index] if investment_index < len(row) else None
            parsed = parse_money(raw_investment)
            if unit and (parsed["amount"] is not None or parsed["label"] not in {None, "-", "Aguardando informação"}):
                values[normalize(unit)] = raw_investment
    workbook.close()
    return values


def load_map(map_path: Path) -> dict[str, Any]:
    with map_path.open(encoding="utf-8") as stream:
        data = json.load(stream)
    return data.get("map", data)


def source_coverage(rows: list[dict[str, Any]], header: str) -> int:
    return sum(1 for row in rows if has_substantive_value(row.get(header)))


def has_substantive_value(value: Any) -> bool:
    text = clean_text(value)
    return bool(text and text not in {"-", "—"} and normalize(text) not in {"aguardandoinformacao", "naoinformado"})


def money_source_coverage(rows: list[dict[str, Any]], header: str) -> int:
    covered = 0
    for row in rows:
        parsed = parse_money(row.get(header))
        if parsed["amount"] is not None or has_substantive_value(parsed["label"]):
            covered += 1
    return covered


def build_payload(workbook_path: Path, map_path: Path) -> dict[str, Any]:
    headers, consolidated_rows = read_rows(workbook_path)
    is_drive_schema = any(row.get("__source_sheet") == DRIVE_SHEET_NAME for row in consolidated_rows)
    formula_errors = scan_drive_formula_errors(workbook_path) if is_drive_schema else []
    municipality_reference = read_municipality_reference(workbook_path)
    # The Drive consolidated already contains every construction row. Re-reading
    # that tab here would duplicate the works and, because its header spans two
    # rows, could also publish a header as if it were a health unit.
    construction_rows = [] if is_drive_schema else read_construction_rows(workbook_path, municipality_reference)
    raw_rows = merge_construction_rows(consolidated_rows, construction_rows)
    source_investments = read_source_investments(workbook_path)
    map_data = load_map(map_path)

    municipality_corrections = {
        "upaengenhovelho": {
            "municipality": "Jaboatão dos Guararapes",
            "reason": "O endereço da linha e a página oficial identificam Jaboatão dos Guararapes; Engenho Velho é o bairro/unidade.",
            "evidenceUrl": "https://portal.saude.pe.gov.br/upas-oss/upa-engenho-velho/",
        },
        "upacurado": {
            "municipality": "Jaboatão dos Guararapes",
            "reason": "O endereço da linha e a página oficial situam a unidade no Curado II, em Jaboatão dos Guararapes.",
            "evidenceUrl": "https://portal.saude.pe.gov.br/upas/upa-curado-medico-fernando-de-lacerda/",
        },
        "upaegranderecife": {
            "municipality": "Abreu e Lima",
            "reason": "O endereço da linha e a página oficial situam a unidade no bairro Desterro, em Abreu e Lima.",
            "evidenceUrl": "https://portal.saude.pe.gov.br/upase-oss/upae-grande-recife-abreu-e-lima/",
        },
        "hospitalmetropolitanonortemiguelarraes": {
            "municipality": "Paulista",
            "reason": "A planilha omite o município; a página oficial informa Paulista no endereço da unidade.",
            "evidenceUrl": "https://portal.saude.pe.gov.br/hospitais-oss/hospital-metropolitano-norte-miguel-arraes-de-alencar/",
        },
    }
    geres_corrections = {
        "hospitalprofagamenonmagalhaes": {
            "geres": "XI",
            "reason": "A unidade está em Serra Talhada, município da XI GERES; a célula I diverge das demais referências territoriais.",
            "evidenceUrl": "https://cartadeservicos.saude.pe.gov.br/endereco-geres/",
        }
    }

    geres_by_municipality: dict[str, list[str]] = defaultdict(list)
    for row in raw_rows:
        source_municipality = clean_text(row.get("NOME MUNICIPIO")) or "Não informado"
        unit_correction = municipality_corrections.get(normalize(row.get("UNIDADE DE SAÚDE")))
        corrected = unit_correction["municipality"] if unit_correction else source_municipality
        geres = clean_text(row.get("GERES"))
        if geres and normalize(row.get("UNIDADE DE SAÚDE")) not in geres_corrections:
            geres_by_municipality[normalize(corrected)].append(geres)

    geres_mode = {
        key: Counter(values).most_common(1)[0][0]
        for key, values in geres_by_municipality.items()
        if values
    }

    units: list[dict[str, Any]] = []
    seen_ids: Counter[str] = Counter()
    corrections: list[dict[str, Any]] = []
    rd_enriched = 0
    geres_enriched = 0
    investments_enriched = 0

    for row in raw_rows:
        unit_name = clean_text(row.get("UNIDADE DE SAÚDE")) or "Unidade sem nome"
        source_municipality = clean_text(row.get("NOME MUNICIPIO")) or "Não informado"
        proposed_correction = municipality_corrections.get(normalize(unit_name))
        correction = proposed_correction if proposed_correction and normalize(source_municipality) != normalize(proposed_correction["municipality"]) else None
        municipality = correction["municipality"] if correction else source_municipality
        if correction:
            corrections.append(
                {
                    "sourceRow": row["__source_row"],
                    "field": "municipality",
                    "sourceValue": source_municipality,
                    "normalizedValue": municipality,
                    "reason": correction["reason"],
                    "evidenceUrl": correction["evidenceUrl"],
                }
            )

        municipality_ref = municipality_reference.get(normalize(municipality))
        source_rd = clean_text(row.get("RD"))
        rd = municipality_ref["rd"] if municipality_ref else source_rd or "Não informado"
        if not source_rd and rd != "Não informado":
            rd_enriched += 1

        source_geres = clean_text(row.get("GERES"))
        geres_correction = geres_corrections.get(normalize(unit_name))
        geres = geres_correction["geres"] if geres_correction else source_geres or geres_mode.get(normalize(municipality)) or "Não informado"
        if geres_correction:
            corrections.append(
                {
                    "sourceRow": row["__source_row"],
                    "field": "geres",
                    "sourceValue": source_geres,
                    "normalizedValue": geres,
                    "reason": geres_correction["reason"],
                    "evidenceUrl": geres_correction["evidenceUrl"],
                }
            )
        if not source_geres and geres != "Não informado":
            geres_enriched += 1

        status = normalize_status(row.get("STATUS"))
        is_construction = status == "Em construção"
        investment_source = row.get("INVESTIMENTO NESTA GESTÃO")
        parsed_investment = parse_money(investment_source)
        works_and_equipment_investment = parse_money(row.get("INVESTIMENTO OBRA E EQUIPAGEM"))
        if parsed_investment["amount"] is None and normalize(unit_name) in source_investments:
            source_sheet_investment = source_investments[normalize(unit_name)]
            replacement = parse_money(source_sheet_investment)
            if replacement["amount"] is not None:
                parsed_investment = replacement
                investments_enriched += 1

        base_id = slugify(unit_name) or f"unidade-{row['__source_row']}"
        seen_ids[base_id] += 1
        unit_id = base_id if seen_ids[base_id] == 1 else f"{base_id}-{seen_ids[base_id]}"

        units.append(
            {
                "id": unit_id,
                "ibgeCode": row.get("__ibge_code") or (municipality_ref["code"] if municipality_ref else None),
                "name": unit_name,
                "municipality": municipality,
                "sourceMunicipality": source_municipality if source_municipality != municipality else None,
                "rd": rd,
                "geres": geres,
                "address": clean_text(row.get("LOCALIZAÇÃO")),
                "status": status,
                "type": normalize_unit_type(row.get("TIPO")),
                "managementType": clean_text(row.get("TIPO GESTÃO")),
                "management": clean_text(row.get("GESTÃO")),
                "beds": None if is_construction else parse_integer(row.get("LEITOS")),
                "plannedBeds": parse_integer(row.get("LEITOS")) if is_construction else None,
                "bedsOpenedInManagement": None if is_construction else parse_integer(row.get("LEITOS ABERTOS NESTA GESTÃO")),
                "openedBedTypes": None if is_construction or not has_substantive_value(row.get("TIPOS DE LEITOS ABERTOS NESTA GESTÃO")) else clean_text(row.get("TIPOS DE LEITOS ABERTOS NESTA GESTÃO")),
                "bedsToOpenAfterRenovation": None if is_construction else parse_integer(row.get("LEITOS A ABRIR COM O FIM DA REFORMA")),
                "profile": clean_text(row.get("PERFIL")),
                "professionals": clean_text(row.get("PROSSIONAIS")),
                "calledProfessionals": parse_integer(row.get("PROFISSIONAIS CONVOCADOS NA GESTÃO")),
                "maintenanceContract": parse_money(row.get("CONTRATO DE MANUTENÇÃO PREDIAL")),
                "managementInvestment": {"amount": None, "label": None} if is_construction else parsed_investment,
                "constructionInvestment": works_and_equipment_investment if is_construction else {"amount": None, "label": None},
                "cofinancing2022": parse_money(row.get("COFINANCIAMENTO 2022")),
                "cofinancing2025": parse_money(row.get("COFINANCIAMENTO 2025")),
                "cofinancingIncreasePercent": parse_percentage(row.get("AUMENTO COFINANCIAMENTO")),
                "osTransfer2025": parse_money(row.get("REPASSE OS 2025")),
                "worksAndEquipmentInvestment": works_and_equipment_investment,
                "professionalBreakdown": {
                    "servers": parse_integer(row.get("PROFISSIONAIS SERVIDORES")),
                    "commissioned": parse_integer(row.get("PROFISSIONAIS COMISSIONADOS")),
                    "clt": parse_integer(row.get("PROFISSIONAIS CLT")),
                    "pj": parse_integer(row.get("PROFISSIONAIS PJ")),
                    "outsourced": parse_integer(row.get("PROFISSIONAIS TERCEIRIZADOS")),
                    "youngApprentices": parse_integer(row.get("JOVEM APRENDIZ")),
                    "temporaryContracts": parse_integer(row.get("CONTRATOS TEMPORÁRIOS")),
                    "secondedFromOtherAgencies": parse_integer(row.get("CEDIDOS DE OUTROS ÓRGÃOS")),
                },
                "mainAdvances": clean_text(row.get("PRINCIPAIS AVANÇOS")),
                "sourceText": clean_text(row.get("TEXTO CONSOLIDADO")),
                "latitude": None,
                "longitude": None,
                "source": {
                    "sheet": row.get("__source_sheet", SHEET_NAME),
                    "row": row["__source_row"],
                    "supplemental": {
                        "sheet": CONSTRUCTION_SHEET_NAME,
                        "row": row["__construction_source_row"],
                    } if row.get("__construction_source_row") else None,
                    "references": row.get("__reference_sources", []),
                    "inheritedFields": row.get("__inherited_fields", []),
                },
                "enrichment": {
                    "rd": "workbook-municipality-template" if municipality_ref else "workbook-row",
                    "geres": "municipality-mode" if not source_geres and geres != "Não informado" else "workbook",
                    "municipalityCorrected": bool(correction),
                    "investment": "construction-sheet" if is_construction and row.get("__construction_source_row") else "source-sheet" if parsed_investment != parse_money(investment_source) else "consolidated",
                    "construction": "construction-sheet" if row.get("__construction_source_row") else "source-row",
                },
            }
        )

    type_counts = Counter(unit["type"] for unit in units)
    status_counts = Counter(unit["status"] for unit in units)
    municipalities = sorted({unit["municipality"] for unit in units}, key=normalize)
    rds = sorted({unit["rd"] for unit in units}, key=normalize)
    geres_values = sorted({unit["geres"] for unit in units}, key=normalize)
    active_units = [unit for unit in units if unit["status"] == "Em funcionamento"]
    construction_units = [unit for unit in units if unit["status"] == "Em construção"]
    units_without_status = [unit for unit in units if unit["status"] == "Não informado"]
    type_counts_by_status = {
        status: dict(sorted(Counter(unit["type"] for unit in units if unit["status"] == status).items()))
        for status in status_counts
    }
    construction_municipalities = {unit["municipality"] for unit in construction_units}
    operational_beds = [unit["beds"] for unit in active_units if unit["beds"] is not None]
    planned_beds = [unit["plannedBeds"] for unit in construction_units if unit["plannedBeds"] is not None]
    non_construction_units = [unit for unit in units if unit["status"] != "Em construção"]
    opened_beds = [unit["bedsOpenedInManagement"] for unit in non_construction_units if unit["bedsOpenedInManagement"] is not None]
    renovation_beds = [unit["bedsToOpenAfterRenovation"] for unit in non_construction_units if unit["bedsToOpenAfterRenovation"] is not None]
    cofinancing_2022 = [unit["cofinancing2022"]["amount"] for unit in units if unit["cofinancing2022"]["amount"] is not None]
    cofinancing_2025 = [unit["cofinancing2025"]["amount"] for unit in units if unit["cofinancing2025"]["amount"] is not None]
    os_transfers_2025 = [unit["osTransfer2025"]["amount"] for unit in units if unit["osTransfer2025"]["amount"] is not None]
    works_investments = [unit["worksAndEquipmentInvestment"]["amount"] for unit in units if unit["worksAndEquipmentInvestment"]["amount"] is not None]
    active_source_rows = [row for row in raw_rows if normalize_status(row.get("STATUS")) == "Em funcionamento"]
    construction_source_rows = [row for row in raw_rows if normalize_status(row.get("STATUS")) == "Em construção"]
    non_construction_source_rows = [row for row in raw_rows if normalize_status(row.get("STATUS")) != "Em construção"]
    identities: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for unit in units:
        identities[(normalize(unit["name"]), normalize(unit["municipality"]))].append(unit)
    possible_duplicates = [
        {
            "unitName": matches[0]["name"],
            "municipality": matches[0]["municipality"],
            "sourceRows": [match["source"]["row"] for match in matches],
        }
        for matches in identities.values()
        if len(matches) > 1
    ]

    quality = {
        "sourceFile": workbook_path.name,
        "sourceSheet": DRIVE_SHEET_NAME if is_drive_schema else SHEET_NAME,
        "sourceSheets": [DRIVE_SHEET_NAME, *DRIVE_SOURCE_SHEETS] if is_drive_schema else [SHEET_NAME, CONSTRUCTION_SHEET_NAME],
        "generatedAt": date.today().isoformat(),
        "sourceSha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "sourceRows": len(raw_rows),
        "constructionRows": len(construction_rows),
        "publishedUnits": len(units),
        "municipalities": len(municipalities),
        "sourceCoverage": {
            "municipality": source_coverage(raw_rows, "NOME MUNICIPIO"),
            "status": source_coverage(raw_rows, "STATUS"),
            "rd": source_coverage(raw_rows, "RD"),
            "geres": source_coverage(raw_rows, "GERES"),
            "address": source_coverage(raw_rows, "LOCALIZAÇÃO"),
            "operationalBeds": sum(1 for row in active_source_rows if parse_integer(row.get("LEITOS")) is not None),
            "plannedBeds": sum(1 for row in construction_source_rows if parse_integer(row.get("LEITOS")) is not None),
            "bedsOpenedInManagement": sum(1 for row in non_construction_source_rows if parse_integer(row.get("LEITOS ABERTOS NESTA GESTÃO")) is not None),
            "openedBedTypes": source_coverage(non_construction_source_rows, "TIPOS DE LEITOS ABERTOS NESTA GESTÃO"),
            "bedsToOpenAfterRenovation": sum(1 for row in non_construction_source_rows if parse_integer(row.get("LEITOS A ABRIR COM O FIM DA REFORMA")) is not None),
            "operationalProfile": source_coverage(active_source_rows, "PERFIL"),
            "plannedProfile": source_coverage(construction_source_rows, "PERFIL"),
            "maintenanceContract": money_source_coverage(active_source_rows, "CONTRATO DE MANUTENÇÃO PREDIAL"),
            "managementInvestment": money_source_coverage(active_source_rows, "INVESTIMENTO NESTA GESTÃO"),
            "constructionInvestment": money_source_coverage(construction_source_rows, "INVESTIMENTO NESTA GESTÃO"),
            "mainAdvances": source_coverage(raw_rows, "PRINCIPAIS AVANÇOS"),
            "cofinancing2022": money_source_coverage(raw_rows, "COFINANCIAMENTO 2022"),
            "cofinancing2025": money_source_coverage(raw_rows, "COFINANCIAMENTO 2025"),
            "cofinancingIncreasePercent": sum(1 for row in raw_rows if parse_percentage(row.get("AUMENTO COFINANCIAMENTO")) is not None),
            "osTransfer2025": money_source_coverage(raw_rows, "REPASSE OS 2025"),
            "worksAndEquipmentInvestment": money_source_coverage(raw_rows, "INVESTIMENTO OBRA E EQUIPAGEM"),
            "professionalBreakdown": sum(1 for unit in units if any(value is not None for value in unit["professionalBreakdown"].values())),
            "coordinates": 0,
        },
        "enrichment": {
            "rdFromWorkbookMunicipalityTemplate": rd_enriched,
            "geresFromMunicipalityMode": geres_enriched,
            "investmentFromSourceSheets": investments_enriched,
            "constructionDetailsFromSheet": sum(1 for row in raw_rows if row.get("__construction_source_row")),
            "documentedCorrections": len(corrections),
            "fieldsInheritedFromPreviousSnapshot": sum(len(row.get("__inherited_fields", [])) for row in raw_rows),
        },
        "formulaErrors": formula_errors,
        "corrections": corrections,
        "possibleDuplicates": possible_duplicates,
        "warnings": [
            "A planilha não contém latitude/longitude; os marcadores do mapa representam o município, não o endereço exato.",
            "Valores ausentes permanecem nulos e são exibidos como Não informado.",
            "RD e código IBGE foram normalizados pela aba 0_Template_Cod_Mun_RD da própria planilha para permitir busca consistente.",
            f"{len(construction_units)} unidades da aba UNIDADES EM CONSTRUÇÃO são publicadas com status Em construção; seus leitos são tratados como previstos e excluídos do total operacional.",
            "Leitos abertos nesta gestão e leitos a abrir após reformas são indicadores de expansão e não são somados ao estoque de leitos em funcionamento nem aos leitos previstos de novas obras.",
            "O rótulo Grandes Emergências na coluna Tipo foi normalizado como Hospital para preservar a taxonomia de unidades do painel.",
            f"{len(units_without_status)} unidades da Rede Credenciada não possuem status na fonte e permanecem como Não informado.",
            f"{len(possible_duplicates)} possível duplicidade por nome e município foi preservada e registrada para validação da área responsável.",
            "O valor Especializado em Tipo Gestão (Hemope) foi preservado, mas requer validação semântica pela área responsável.",
            *([f"{len(formula_errors)} erro(s) de fórmula foram encontrados no consolidado do Drive; campos equivalentes foram recuperados das abas de origem quando disponíveis."] if formula_errors else []),
            *(["Campos ausentes no consolidado são complementados pelas abas de origem e, apenas quando necessário, pelo último JSON válido; a procedência fica registrada por unidade."] if is_drive_schema else []),
        ],
        "headers": headers,
    }

    return {
        "meta": {
            "title": "Unidades de Saúde de Pernambuco",
            "sourceLabel": "Google Planilhas · sincronização automática" if is_drive_schema else "Planilha Saúde · 20 ago 2026",
            "generatedAt": quality["generatedAt"],
            "sourceUrl": "https://docs.google.com/spreadsheets/d/1I0rhyE5D9jvFqe6IHPsyCxSyr3QZEL-AgL4p2DMGTEE/edit" if is_drive_schema else None,
            "totalUnits": len(units),
            "activeUnits": len(active_units),
            "constructionUnits": len(construction_units),
            "unitsWithoutStatus": len(units_without_status),
            "constructionMunicipalities": len(construction_municipalities),
            "totalMunicipalities": len(municipalities),
            "totalBeds": sum(operational_beds),
            "unitsWithBeds": len(operational_beds),
            "plannedBeds": sum(planned_beds),
            "constructionUnitsWithBeds": len(planned_beds),
            "bedsOpenedInManagement": sum(opened_beds),
            "unitsWithBedsOpenedInManagement": len(opened_beds),
            "bedsToOpenAfterRenovation": sum(renovation_beds),
            "unitsWithBedsToOpenAfterRenovation": len(renovation_beds),
            "cofinancing2022": sum(cofinancing_2022),
            "unitsWithCofinancing2022": len(cofinancing_2022),
            "cofinancing2025": sum(cofinancing_2025),
            "unitsWithCofinancing2025": len(cofinancing_2025),
            "osTransfers2025": sum(os_transfers_2025),
            "unitsWithOsTransfers2025": len(os_transfers_2025),
            "worksAndEquipmentInvestment": sum(works_investments),
            "unitsWithWorksAndEquipmentInvestment": len(works_investments),
            "typeCounts": dict(sorted(type_counts.items())),
            "typeCountsByStatus": type_counts_by_status,
            "statusCounts": dict(status_counts),
        },
        "filters": {
            "municipalities": municipalities,
            "rds": rds,
            "geres": geres_values,
            "types": sorted(type_counts, key=normalize),
            "statuses": [status for status in ("Em funcionamento", "Em construção", "Não informado") if status in status_counts],
        },
        "units": units,
        "map": {
            **map_data,
            "fallbackCenters": {
                "2605459": {"x": 935, "y": 48, "label": "Fernando de Noronha"}
            },
        },
        "dataQuality": quality,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Gera o JSON do painel a partir das abas Consolidado e UNIDADES EM CONSTRUÇÃO.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--map-data", type=Path, default=DEFAULT_MAP_DATA)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--quality-output", type=Path, default=DEFAULT_QUALITY_OUTPUT)
    args = parser.parse_args()

    payload = build_payload(args.workbook, args.map_data)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.quality_output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    args.quality_output.write_text(json.dumps(payload["dataQuality"], ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        json.dumps(
            {
                "output": str(args.output),
                "units": payload["meta"]["totalUnits"],
                "activeUnits": payload["meta"]["activeUnits"],
                "constructionUnits": payload["meta"]["constructionUnits"],
                "municipalities": payload["meta"]["totalMunicipalities"],
                "operationalBeds": payload["meta"]["totalBeds"],
                "plannedBeds": payload["meta"]["plannedBeds"],
                "bedsOpenedInManagement": payload["meta"]["bedsOpenedInManagement"],
                "bedsToOpenAfterRenovation": payload["meta"]["bedsToOpenAfterRenovation"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
